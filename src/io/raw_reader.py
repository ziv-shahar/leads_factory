"""Raw file reader with local filesystem support (S3-ready interface)."""
import os
from pathlib import Path
from typing import List, Tuple, Optional, Dict
from bs4 import BeautifulSoup
import json
import fnmatch

# Document parsing libraries
try:
    from pypdf import PdfReader
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from openpyxl import load_workbook
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False


class RawFileReader:
    """Read raw files from local filesystem (can be swapped to S3 later)."""

    def __init__(self, bucket_path: str):
        """
        Initialize reader.

        Args:
            bucket_path: Path to local bucket directory (or S3 bucket name in future)
        """
        self.bucket_path = Path(bucket_path)

    def list_files(self, pattern: str = "*") -> List[str]:
        """
        List files in bucket matching pattern.

        Args:
            pattern: Glob pattern (default: all files)

        Returns:
            List of file paths
        """
        if not self.bucket_path.exists():
            return []

        files = []
        for path in self.bucket_path.glob(pattern):
            if path.is_file():
                files.append(str(path))

        return sorted(files)

    def read_file(self, file_path: str) -> Tuple[str, str]:
        """
        Read and parse file content.

        Args:
            file_path: Path to file

        Returns:
            Tuple of (content, file_type)
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        # Determine file type
        suffix = path.suffix.lower()

        # Read content based on file type
        if suffix in ['.html', '.htm']:
            return self._read_html(path), 'html'
        elif suffix == '.json':
            return self._read_json(path), 'json'
        elif suffix in ['.txt', '.text']:
            return self._read_text(path), 'text'
        elif suffix == '.pdf':
            return self._read_pdf(path), 'pdf'
        elif suffix in ['.docx', '.doc']:
            return self._read_docx(path), 'docx'
        elif suffix in ['.xlsx', '.xls']:
            return self._read_xlsx(path), 'xlsx'
        else:
            # Fallback to text
            return self._read_text(path), 'text'

    def _read_html(self, path: Path) -> str:
        """Read and extract text from HTML."""
        with open(path, 'r', encoding='utf-8') as f:
            html_content = f.read()

        # Parse HTML and extract text
        soup = BeautifulSoup(html_content, 'html.parser')

        # Remove script and style elements
        for script in soup(["script", "style"]):
            script.decompose()

        # Get text
        text = soup.get_text(separator='\n')

        # Clean up whitespace
        lines = (line.strip() for line in text.splitlines())
        chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
        text = '\n'.join(chunk for chunk in chunks if chunk)

        return text

    def _read_json(self, path: Path) -> str:
        """Read JSON and convert to readable format."""
        with open(path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        # Pretty print JSON for LLM readability
        return json.dumps(data, indent=2)

    def _read_text(self, path: Path) -> str:
        """Read plain text file."""
        with open(path, 'r', encoding='utf-8') as f:
            return f.read()

    def _read_pdf(self, path: Path) -> str:
        """Read and extract text from PDF."""
        if not PDF_AVAILABLE:
            raise ImportError(
                "pypdf is not installed. Install it with: pip install pypdf"
            )

        try:
            reader = PdfReader(str(path))
            text_parts = []

            # Extract text from each page
            for page_num, page in enumerate(reader.pages, 1):
                text = page.extract_text()
                if text.strip():
                    text_parts.append(f"=== PAGE {page_num} ===\n{text}")

            if not text_parts:
                return "[PDF contains no extractable text - may be scanned/image-based]"

            return "\n\n".join(text_parts)

        except Exception as e:
            return f"[Error reading PDF: {str(e)}]"

    def _read_docx(self, path: Path) -> str:
        """Read and extract text from DOCX."""
        if not DOCX_AVAILABLE:
            raise ImportError(
                "python-docx is not installed. Install it with: pip install python-docx"
            )

        try:
            doc = Document(str(path))
            text_parts = []

            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # Extract tables
            if doc.tables:
                text_parts.append("\n=== TABLES ===")
                for table_num, table in enumerate(doc.tables, 1):
                    text_parts.append(f"\nTable {table_num}:")
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        if row_text.strip():
                            text_parts.append(row_text)

            if not text_parts:
                return "[DOCX contains no extractable text]"

            return "\n".join(text_parts)

        except Exception as e:
            return f"[Error reading DOCX: {str(e)}]"

    def _read_xlsx(self, path: Path) -> str:
        """Read and extract data from XLSX as formatted text."""
        if not XLSX_AVAILABLE:
            raise ImportError(
                "openpyxl is not installed. Install it with: pip install openpyxl"
            )

        try:
            workbook = load_workbook(str(path), read_only=True, data_only=True)
            text_parts = []

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_parts.append(f"=== SHEET: {sheet_name} ===\n")

                # Get all rows with data
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    # Skip empty rows
                    if any(cell is not None and str(cell).strip() for cell in row):
                        row_text = " | ".join(
                            str(cell) if cell is not None else ""
                            for cell in row
                        )
                        rows.append(row_text)

                if rows:
                    # Limit to first 100 rows per sheet for LLM context
                    if len(rows) > 100:
                        text_parts.append("\n".join(rows[:100]))
                        text_parts.append(f"\n[... {len(rows) - 100} more rows omitted ...]")
                    else:
                        text_parts.append("\n".join(rows))
                else:
                    text_parts.append("[Empty sheet]")

                text_parts.append("")  # Blank line between sheets

            workbook.close()

            if not text_parts:
                return "[XLSX contains no data]"

            return "\n".join(text_parts)

        except Exception as e:
            return f"[Error reading XLSX: {str(e)}]"

    def check_merge_config(self, directory: Path) -> Optional[List[str]]:
        """
        Check if directory has merge.txt configuration.

        Args:
            directory: Directory to check

        Returns:
            List of glob patterns to merge, or None if no merge.txt exists
        """
        merge_file = directory / "merge.txt"
        if not merge_file.exists():
            return None

        patterns = []
        try:
            with open(merge_file, 'r', encoding='utf-8') as f:
                for line in f:
                    line = line.strip()
                    # Skip empty lines and comments
                    if line and not line.startswith('#'):
                        patterns.append(line)
        except Exception as e:
            print(f"  ⚠ Error reading merge.txt: {str(e)}")
            return None

        return patterns if patterns else None

    def should_merge_file(self, file_name: str, patterns: List[str]) -> bool:
        """
        Check if file matches merge patterns.

        Args:
            file_name: Name of file to check
            patterns: List of glob patterns (supports ! for exclusion)

        Returns:
            True if file should be included in merge
        """
        # Start with not included
        included = False

        for pattern in patterns:
            # Handle exclusion patterns (!)
            if pattern.startswith('!'):
                exclude_pattern = pattern[1:]
                if fnmatch.fnmatch(file_name, exclude_pattern):
                    return False  # Explicitly excluded
            else:
                # Inclusion pattern
                if fnmatch.fnmatch(file_name, pattern):
                    included = True

        return included

    def merge_directory_files(self, directory: Path, patterns: List[str]) -> Tuple[str, List[str]]:
        """
        Merge multiple files from a directory based on patterns.

        Args:
            directory: Directory containing files to merge
            patterns: List of glob patterns

        Returns:
            Tuple of (merged_content, list_of_merged_files)
        """
        merged_parts = []
        merged_files = []

        # Get all files in directory
        files = sorted([f for f in directory.iterdir() if f.is_file()])

        for file_path in files:
            file_name = file_path.name

            # Skip merge.txt itself
            if file_name == "merge.txt":
                continue

            # Check if file matches patterns
            if self.should_merge_file(file_name, patterns):
                try:
                    # Read file content
                    content, file_type = self.read_file(str(file_path))

                    # Add to merged content with clear separator
                    merged_parts.append(f"{'=' * 80}")
                    merged_parts.append(f"FILE: {file_name}")
                    merged_parts.append(f"TYPE: {file_type}")
                    merged_parts.append(f"{'=' * 80}")
                    merged_parts.append(content)
                    merged_parts.append("")  # Blank line between files

                    merged_files.append(str(file_path))

                except Exception as e:
                    # If a file fails, add error note but continue
                    merged_parts.append(f"{'=' * 80}")
                    merged_parts.append(f"FILE: {file_name}")
                    merged_parts.append(f"ERROR: {str(e)}")
                    merged_parts.append(f"{'=' * 80}")
                    merged_parts.append("")

        if not merged_parts:
            return "[No files matched merge patterns]", []

        return "\n".join(merged_parts), merged_files

    def list_files_with_merge(self, pattern: str = "*") -> List[Dict[str, any]]:
        """
        List files with merge support - non-recursive, explicit discovery.

        Discovery rules (processes only immediate children of RAW_DATA_BUCKET):
        1. Files at root level: Process individually
        2. Directory with merge.txt: Process as merged entry (combines files per merge.txt patterns)
        3. Directory with files (no merge.txt): Process each file individually
        4. Directory with only subdirectories (no direct files): Ignore
        5. Does NOT recurse into nested directories automatically

        Examples:
          raw_data_bucket/
            article.html              → Processed as file
            news/
              story1.html             → Processed as file
              story2.html             → Processed as file
            contracts/
              merge.txt               → Directory processed as merged entry
              doc1.pdf
              doc2.docx
            reports/
              2024/
                data.xlsx             → Ignored (nested)

        Returns a list of file entries. Each entry is either:
        - A single file: {"type": "file", "path": "...", "files": ["..."]}
        - A merged directory: {"type": "merged", "path": "dir_path", "files": ["file1", "file2", ...]}

        Args:
            pattern: Glob pattern (currently unused, kept for API compatibility)

        Returns:
            List of file entry dicts
        """
        if not self.bucket_path.exists():
            return []

        entries = []

        # Iterate through immediate children of bucket_path
        for item in sorted(self.bucket_path.iterdir()):
            if item.is_file():
                # Individual file at root level - add it
                entries.append({
                    "type": "file",
                    "path": str(item),
                    "files": [str(item)]
                })

            elif item.is_dir():
                # Check if directory has merge.txt
                merge_patterns = self.check_merge_config(item)

                if merge_patterns:
                    # Directory with merge.txt - create merged entry
                    _, merged_files = self.merge_directory_files(item, merge_patterns)

                    if merged_files:
                        entries.append({
                            "type": "merged",
                            "path": str(item),
                            "files": merged_files,
                            "merge_patterns": merge_patterns
                        })
                else:
                    # Directory without merge.txt - check if it has files directly in it
                    has_files = any(f.is_file() for f in item.iterdir())

                    if has_files:
                        # Process each file in this directory individually
                        for file_path in sorted(item.iterdir()):
                            if file_path.is_file():
                                entries.append({
                                    "type": "file",
                                    "path": str(file_path),
                                    "files": [str(file_path)]
                                })
                    # Else: directory with only subdirectories - ignore it

        return entries

    def get_relative_path(self, file_path: str) -> str:
        """Get path relative to bucket root."""
        path = Path(file_path)
        try:
            return str(path.relative_to(self.bucket_path))
        except ValueError:
            return str(path)
